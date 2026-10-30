"""
Tests for STEP 4: Output Modes (Mode A Editable, Mode B Exact Layout, and Auto Dispatcher).
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.analyzer import analyze_pdf
from app.docx_modes import build_docx, build_mode_a_editable, build_mode_b_exact
from app.extractor import extract_pdf_document
from app.font_resolver import FontResolver
from tests.conftest import create_khmer_text_pdf


def test_build_mode_a_editable(tmp_path: Path):
    pdf_path = create_khmer_text_pdf(tmp_path / "test_a.pdf")
    docx_path = tmp_path / "out_mode_a.docx"

    analysis = analyze_pdf(pdf_path)
    resolver = FontResolver()
    res = build_mode_a_editable(pdf_path, docx_path, analysis, resolver)

    assert res.exists()
    assert res.stat().st_size > 0

    doc = Document(str(docx_path))
    assert len(doc.paragraphs) > 0
    # Check section size
    assert doc.sections[0].page_width.pt > 500


def test_build_mode_b_exact(tmp_path: Path):
    pdf_path = create_khmer_text_pdf(tmp_path / "test_b.pdf")
    docx_path = tmp_path / "out_mode_b.docx"

    analysis = analyze_pdf(pdf_path)
    pages = extract_pdf_document(pdf_path)
    resolver = FontResolver()
    res = build_mode_b_exact(pdf_path, docx_path, analysis, pages, resolver)

    assert res.exists()
    assert res.stat().st_size > 0

    doc = Document(str(docx_path))
    assert len(doc.paragraphs) > 0

    # In Mode B, paragraphs have w:framePr attached
    found_frame_pr = False
    for p in doc.paragraphs:
        pPr = p._p.get_or_add_pPr()
        framePr = pPr.find(qn("w:framePr"))
        if framePr is not None:
            found_frame_pr = True
            assert framePr.get(qn("w:hAnchor")) == "page"
            assert framePr.get(qn("w:vAnchor")) == "page"
            break

    assert found_frame_pr is True


def test_build_docx_hybrid(tmp_path: Path):
    pdf_path = create_khmer_text_pdf(tmp_path / "test_hybrid.pdf")
    out_a = tmp_path / "hybrid_a.docx"
    out_b = tmp_path / "hybrid_b.docx"

    res_a = build_docx(pdf_path, out_a, mode="editable")
    assert res_a.exists()

    res_b = build_docx(pdf_path, out_b, mode="exact")
    assert res_b.exists()

