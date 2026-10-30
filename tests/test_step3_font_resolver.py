"""
Tests for STEP 3: Dynamic Font Resolution, Font Embedding, and Text Fitting.
"""

from __future__ import annotations

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.font_resolver import (
    FontResolver,
    apply_khmer_run_properties,
    obfuscate_font_data,
)


def test_font_name_resolution():
    resolver = FontResolver()

    # Subset prefix stripping + mapping
    res1 = resolver.resolve_font("ABCDEF+KhmerOSBattambang-Bold")
    assert res1.resolved_name == "Khmer OS Battambang"
    assert res1.resolution_source in ("map", "system")

    # Muol Light mapping
    res2 = resolver.resolve_font("XYZABC+KhmerOSMuolLight")
    assert res2.resolved_name == "Khmer OS Muol Light"

    # Fallback when completely unknown font
    res3 = resolver.resolve_font("SomeRandomNonExistentFont12345")
    assert res3.resolved_name in ["Noto Sans Khmer", "Khmer UI", "Leelawadee UI", "Khmer OS Battambang"]


def test_text_fitting():
    resolver = FontResolver()
    text = "ព្រះរាជាណាចក្រកម្ពុជា"

    # Fits naturally when target width is huge
    fit1 = resolver.fit_text(text, "Khmer OS Battambang", 12.0, 500.0)
    assert fit1.adjusted_size == 12.0
    assert fit1.spacing_twips == 0

    # Condenses and reduces by at most 8% when target width is cramped
    fit2 = resolver.fit_text(text, "Khmer OS Battambang", 12.0, 30.0)
    assert fit2.spacing_twips < 0  # condensed spacing
    assert fit2.adjusted_size >= 12.0 * 0.92  # at most 8% font reduction


def test_font_obfuscation():
    sample_font = b"\x00\x01\x00\x00" + b"\x42" * 60
    guid = "{A38B8682-1D5B-4C91-9DEE-2F3B66D68E56}"
    obf = obfuscate_font_data(sample_font, guid)
    assert len(obf) == len(sample_font)
    assert obf != sample_font

    # Re-obfuscation with same key restores original bytes
    deobf = obfuscate_font_data(obf, guid)
    assert deobf == sample_font


def test_apply_khmer_run_properties():
    doc = Document()
    para = doc.add_paragraph()
    run = para.add_run("សួស្តីកម្ពុជា")

    apply_khmer_run_properties(
        run,
        font_name="Khmer OS Muol Light",
        font_size_pt=14.0,
        is_bold=True,
        is_italic=False,
        color_hex="#1E3A8A",
        spacing_twips=-10,
    )

    rPr = run._r.get_or_add_rPr()

    # w:rFonts
    rFonts = rPr.find(qn("w:rFonts"))
    assert rFonts is not None
    assert rFonts.get(qn("w:cs")) == "Khmer OS Muol Light"
    assert rFonts.get(qn("w:ascii")) == "Khmer OS Muol Light"

    # w:sz and w:szCs (14pt * 2 = 28 half-points)
    assert rPr.find(qn("w:sz")).get(qn("w:val")) == "28"
    assert rPr.find(qn("w:szCs")).get(qn("w:val")) == "28"

    # w:b and w:bCs
    assert rPr.find(qn("w:b")) is not None
    assert rPr.find(qn("w:bCs")) is not None

    # w:lang w:bidi
    lang = rPr.find(qn("w:lang"))
    assert lang is not None
    assert lang.get(qn("w:bidi")) == "km-KH"

    # w:spacing
    sp = rPr.find(qn("w:spacing"))
    assert sp is not None
    assert sp.get(qn("w:val")) == "-10"

    # w:color
    col = rPr.find(qn("w:color"))
    assert col is not None
    assert col.get(qn("w:val")) == "1E3A8A"

